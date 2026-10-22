#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
xlsx 转 JSON 转换工具（增强版）
将 Excel 文件转换为紧凑的 JSON 格式，并生成预计算摘要数据。

功能：
1. 原始明细 JSON（数组形式）- 向后兼容
2. 单周预计算摘要 - 前端直接加载，跳过聚合计算
3. 双周环比对比摘要 - 前端直接渲染，跳过环比计算

使用方法：
  python convert_xlsx_to_json.py <xlsx文件路径>              # 单文件转换
  python convert_xlsx_to_json.py --batch <data目录路径>       # 批处理 + 生成对比摘要
"""

import sys
import os
import json
from openpyxl import load_workbook


# ========== 字段映射配置 ==========
SUBLET_FIELDS = {
    'date': 'part_day(day)',
    'month': 'part_day(month)',
    'channel': '三方渠道',
    'sublessorId': 'pid',
    'sublessorName': 'pusername',
    'gameName': 'game_name',
    'gameService': 'game_service',
    'shopId': 'shop_id',
    'shopName': 'shop_name',
    'isSelfOperated': '是否自营',
    'isAutoOrder': '是否自动化订单',
    'orderCount': '订单量',
    'settledOrderCount': '结算订单量',
    'cancelCount': '撤单量',
    'orderAmount': '订单金额',
    'settledAmount': '结算订单金额',
    'sublessorProfit': '转租商收益',
    'platformProfit': '平台收益（含服务费）',
    'thirdPartyOrders': '三方订单量',
    'thirdPartyShipments': '三方发货量',
    'thirdPartyAmount': '三方订单金额'
}

DIST_FIELDS = {
    'date': 'part_day(day)',
    'distributorId': 'user_id',
    'distributorType': 'add_from',
    'gameName': 'game_name',
    'orderPortal': '下单端口',
    'isSelfOperated': '是否自营',
    'orderCount': '订单量',
    'normalOrderCount': '正常订单量',
    'settledOrderCount': '正常结算订单量',
    'cancelCount': '撤单量',
    'orderAmount': '订单金额',
    'userPayment': '用户实付',
    'platformProfit': '平台收益',
    'platformMarkup': '平台加价',
    'distributorProfit': '分销收益',
    'distributorMarkup': '分销加价',
    'ownerProfit': '号主收益',
    'selfOperatedOrders': '自营结算订单量',
    'nonSelfOperatedOrders': '非自营结算订单量'
}


# ========== 工具函数 ==========
def safe_num(val):
    if val is None or val == '':
        return 0
    try:
        return float(val)
    except (ValueError, TypeError):
        return 0


def shop_agg_key(sublessor_id, sublessor_name, shop_id, game_name='', channel=None):
    """构建店铺聚合的唯一key"""
    sub_id = sublessor_id or ('name:' + (sublessor_name or '未知转租商'))
    parts = [str(sub_id), str(shop_id or '')]
    if game_name:
        parts.append(str(game_name))
    if channel is not None and channel != '':
        parts.append(str(channel))
    return '||'.join(parts)


def calc_change_rate(cur, prev):
    if prev == 0:
        if cur == 0:
            return 0
        return 1 if cur > 0 else -1
    return (cur - prev) / abs(prev)


def get_date_range(rows, date_field):
    if not rows:
        return {'min': None, 'max': None}
    min_val = None
    max_val = None
    for row in rows:
        d = row.get(date_field)
        if d:
            ds = str(d)
            if min_val is None or ds < min_val:
                min_val = ds
            if max_val is None or ds > max_val:
                max_val = ds
    return {'min': min_val, 'max': max_val}


# ========== 数据类型检测 ==========
def detect_data_type(headers):
    if not headers:
        return 'unknown'
    header_set = set(str(h).strip() for h in headers)

    sublet_strong = ['三方渠道', '转租商收益', 'shop_id', 'shop_name', 'pid', 'pusername',
                     '三方订单量', '三方发货量', '三方订单金额']
    sublet_strong_score = sum(1 for m in sublet_strong if m in header_set)

    sublet_weak = ['结算订单金额', '平台收益（含服务费）', 'quick_type']
    sublet_weak_score = sum(1 for m in sublet_weak if m in header_set)

    sublet_score = sublet_strong_score * 2 + sublet_weak_score

    dist_strong = ['分销收益', '分销加价', 'add_from', '号主收益', '平台加价']
    dist_strong_score = sum(1 for m in dist_strong if m in header_set)

    dist_weak = ['用户实付', '下单端口', '自营结算订单量', '非自营结算订单量', '正常结算订单量']
    dist_weak_score = sum(1 for m in dist_weak if m in header_set)

    dist_score = dist_strong_score * 2 + dist_weak_score

    if sublet_score > dist_score and sublet_strong_score >= 1:
        return 'sublet'
    if dist_score > sublet_score and dist_strong_score >= 1:
        return 'distribution'
    if sublet_strong_score > dist_strong_score and sublet_strong_score >= 2:
        return 'sublet'
    if dist_strong_score > sublet_strong_score and dist_strong_score >= 2:
        return 'distribution'
    return 'unknown'


# ========== 转租数据处理 ==========
def process_sublet_data(raw_rows, headers):
    f = SUBLET_FIELDS
    header_set = set(str(h).strip() for h in headers)

    # 字段兼容性处理
    date_field = f['date'] if f['date'] in header_set else (
        'day' if 'day' in header_set else ('日期' if '日期' in header_set else f['date']))
    month_field = f['month'] if f['month'] in header_set else None

    channel_field = f['channel'] if f['channel'] in header_set else (
        '渠道' if '渠道' in header_set else ('channel' if 'channel' in header_set else None))

    sublessor_id_field = f['sublessorId'] if f['sublessorId'] in header_set else (
        '转租商ID' if '转租商ID' in header_set else None)

    sublessor_name_field = f['sublessorName'] if f['sublessorName'] in header_set else (
        '转租商名' if '转租商名' in header_set else ('pname' if 'pname' in header_set else None))

    shop_id_field = f['shopId'] if f['shopId'] in header_set else (
        '店铺ID' if '店铺ID' in header_set else f['shopId'])

    shop_name_field = f['shopName'] if f['shopName'] in header_set else (
        '店铺名' if '店铺名' in header_set else f['shopName'])

    settled_order_field = f['settledOrderCount'] if f['settledOrderCount'] in header_set else (
        '正常结算订单量' if '正常结算订单量' in header_set else (
            '正常订单量' if '正常订单量' in header_set else f['orderCount']))

    settled_amount_field = f['settledAmount'] if f['settledAmount'] in header_set else (
        '结算金额' if '结算金额' in header_set else f['settledAmount'])

    platform_profit_field = f['platformProfit'] if f['platformProfit'] in header_set else (
        '平台收益' if '平台收益' in header_set else f['platformProfit'])

    # 构建行索引映射
    header_idx = {h: i for i, h in enumerate(headers)}

    def get_val(row, field_name):
        idx = header_idx.get(field_name)
        if idx is not None and idx < len(row):
            return row[idx]
        return ''

    # 检查是否有渠道数据
    has_channel_data = False
    if channel_field:
        for row in raw_rows[:100]:
            val = get_val(row, channel_field)
            if val and str(val).strip() != '':
                has_channel_data = True
                break

    # 标准化字段 - 一次遍历完成所有聚合
    rows = []
    totals = {
        'orderCount': 0, 'settledOrderCount': 0, 'cancelCount': 0,
        'orderAmount': 0, 'settledAmount': 0, 'platformProfit': 0,
        'sublessorProfit': 0, 'thirdPartyOrders': 0, 'thirdPartyShipments': 0,
        'thirdPartyAmount': 0
    }
    channel_map = {}
    game_map = {}
    shop_map = {}
    game_shop_map = {}
    xianyu_detail_map = {}
    taobao_detail_map = {}

    for row in raw_rows:
        date_val = get_val(row, date_field)
        month_val = get_val(row, month_field) if month_field else (
            str(date_val)[:6] if date_val else '')
        channel = get_val(row, channel_field) if channel_field else '未知渠道'
        channel = channel if channel and str(channel).strip() else '未知渠道'

        sublessor_id = get_val(row, sublessor_id_field) if sublessor_id_field else ''
        sublessor_name = get_val(row, sublessor_name_field) if sublessor_name_field else ''
        game_name = get_val(row, f['gameName']) or '未知游戏'
        game_service = get_val(row, f['gameService']) or ''
        shop_id = str(get_val(row, shop_id_field) or '')
        shop_name = get_val(row, shop_name_field) or ''
        is_self = get_val(row, f['isSelfOperated']) or ''
        is_auto = get_val(row, f['isAutoOrder']) or ''
        order_count = safe_num(get_val(row, f['orderCount']))
        settled_order = safe_num(get_val(row, settled_order_field))
        cancel_count = safe_num(get_val(row, f['cancelCount']))
        order_amount = safe_num(get_val(row, f['orderAmount']))
        settled_amount = safe_num(get_val(row, settled_amount_field))
        sublessor_profit = safe_num(get_val(row, f['sublessorProfit']))
        platform_profit = safe_num(get_val(row, platform_profit_field))
        third_orders = safe_num(get_val(row, f['thirdPartyOrders']))
        third_shipments = safe_num(get_val(row, f['thirdPartyShipments']))
        third_amount = safe_num(get_val(row, f['thirdPartyAmount']))

        r = {
            'date': date_val,
            'month': month_val,
            'channel': channel,
            'sublessorId': sublessor_id,
            'sublessorName': sublessor_name,
            'gameName': game_name,
            'gameService': game_service,
            'shopId': shop_id,
            'shopName': shop_name,
            'isSelfOperated': is_self,
            'isAutoOrder': is_auto,
            'orderCount': order_count,
            'settledOrderCount': settled_order,
            'cancelCount': cancel_count,
            'orderAmount': order_amount,
            'settledAmount': settled_amount,
            'sublessorProfit': sublessor_profit,
            'platformProfit': platform_profit,
            'thirdPartyOrders': third_orders,
            'thirdPartyShipments': third_shipments,
            'thirdPartyAmount': third_amount
        }
        rows.append(r)

        # 实时聚合 totals
        totals['orderCount'] += order_count
        totals['settledOrderCount'] += settled_order
        totals['cancelCount'] += cancel_count
        totals['orderAmount'] += order_amount
        totals['settledAmount'] += settled_amount
        totals['platformProfit'] += platform_profit
        totals['sublessorProfit'] += sublessor_profit
        totals['thirdPartyOrders'] += third_orders
        totals['thirdPartyShipments'] += third_shipments
        totals['thirdPartyAmount'] += third_amount

        # 实时聚合渠道
        ch = channel
        if ch not in channel_map:
            channel_map[ch] = {
                'channel': ch, 'settledOrderCount': 0, 'platformProfit': 0,
                'settledAmount': 0, 'cancelCount': 0, 'orderCount': 0
            }
        channel_map[ch]['settledOrderCount'] += settled_order
        channel_map[ch]['platformProfit'] += platform_profit
        channel_map[ch]['settledAmount'] += settled_amount
        channel_map[ch]['cancelCount'] += cancel_count
        channel_map[ch]['orderCount'] += order_count

        # 实时聚合游戏
        g = game_name
        if g not in game_map:
            game_map[g] = {
                'gameName': g, 'settledOrderCount': 0, 'settledAmount': 0,
                'platformProfit': 0, 'cancelCount': 0, 'orderCount': 0
            }
        game_map[g]['settledOrderCount'] += settled_order
        game_map[g]['settledAmount'] += settled_amount
        game_map[g]['platformProfit'] += platform_profit
        game_map[g]['cancelCount'] += cancel_count
        game_map[g]['orderCount'] += order_count

        # 实时聚合店铺（按转租商+店铺+游戏，跨渠道汇总）
        key = shop_agg_key(sublessor_id, sublessor_name, shop_id, game_name)
        if key not in shop_map:
            shop_map[key] = {
                'shopId': shop_id,
                'shopName': shop_name or ('手动开单' if not shop_id else ''),
                'sublessorName': sublessor_name,
                'sublessorId': sublessor_id,
                'gameName': game_name,
                'channel': channel,
                'settledOrderCount': 0, 'settledAmount': 0,
                'platformProfit': 0, 'sublessorProfit': 0,
                'cancelCount': 0, 'orderCount': 0
            }
        shop_map[key]['settledOrderCount'] += settled_order
        shop_map[key]['settledAmount'] += settled_amount
        shop_map[key]['platformProfit'] += platform_profit
        shop_map[key]['sublessorProfit'] += sublessor_profit
        shop_map[key]['cancelCount'] += cancel_count
        shop_map[key]['orderCount'] += order_count

        # 实时聚合游戏→店铺（用于异常预警下钻）
        if game_name not in game_shop_map:
            game_shop_map[game_name] = {}
        gs_key = shop_agg_key(sublessor_id, sublessor_name, shop_id)
        if gs_key not in game_shop_map[game_name]:
            game_shop_map[game_name][gs_key] = {
                'shopId': shop_id,
                'shopName': shop_name or ('手动开单' if not shop_id else ''),
                'sublessorId': sublessor_id,
                'sublessorName': sublessor_name,
                'settledOrderCount': 0, 'settledAmount': 0
            }
        game_shop_map[game_name][gs_key]['settledOrderCount'] += settled_order
        game_shop_map[game_name][gs_key]['settledAmount'] += settled_amount

        # 闲鱼渠道细分
        if channel == '闲鱼':
            xt = '闲鱼'
            if xt not in xianyu_detail_map:
                xianyu_detail_map[xt] = {'type': xt, 'settledOrderCount': 0, 'settledAmount': 0, 'platformProfit': 0}
            xianyu_detail_map[xt]['settledOrderCount'] += settled_order
            xianyu_detail_map[xt]['settledAmount'] += settled_amount
            xianyu_detail_map[xt]['platformProfit'] += platform_profit

        # 淘宝渠道细分
        if channel in ('淘宝', '淘宝中心场'):
            tt = channel
            if tt not in taobao_detail_map:
                taobao_detail_map[tt] = {'type': tt, 'settledOrderCount': 0, 'settledAmount': 0, 'platformProfit': 0}
            taobao_detail_map[tt]['settledOrderCount'] += settled_order
            taobao_detail_map[tt]['settledAmount'] += settled_amount
            taobao_detail_map[tt]['platformProfit'] += platform_profit

    # 计算 totals 派生值
    totals['cancelRate'] = totals['cancelCount'] / totals['orderCount'] if totals['orderCount'] > 0 else 0
    totals['shipSuccessRate'] = (
        min(1, totals['thirdPartyShipments'] / totals['thirdPartyOrders'])
        if totals['thirdPartyOrders'] > 0 else 0
    )

    date_range = get_date_range(rows, 'date')

    # 渠道数据排序
    channel_data = sorted(channel_map.values(), key=lambda x: x['settledOrderCount'], reverse=True)

    xianyu_detail = list(xianyu_detail_map.values())
    taobao_detail = list(taobao_detail_map.values())

    # 游戏数据 + 撤单率
    game_data = list(game_map.values())
    for g in game_data:
        g['cancelRate'] = g['cancelCount'] / g['orderCount'] if g['orderCount'] > 0 else 0

    # TOP20
    top20_by_amount = sorted(game_data, key=lambda x: x['settledAmount'], reverse=True)[:20]
    top20_by_orders = sorted(game_data, key=lambda x: x['settledOrderCount'], reverse=True)[:20]

    # TOP5游戏平均撤单率
    top5_games = sorted(game_data, key=lambda x: x['settledOrderCount'], reverse=True)[:5]
    top5_avg_cancel_rate = (
        sum(g['cancelRate'] for g in top5_games) / len(top5_games)
        if top5_games else 0
    )

    # 店铺撤单率
    shop_data = list(shop_map.values())
    for s in shop_data:
        s['cancelRate'] = s['cancelCount'] / s['orderCount'] if s['orderCount'] > 0 else 0

    # 商户-游戏撤单率预警
    game_cancel_map = {g['gameName']: g['cancelRate'] for g in game_data}
    merchant_cancel_warning = []
    for s in shop_data:
        game_rate = game_cancel_map.get(s['gameName'], 0)
        if s['orderCount'] >= 10 and s['cancelRate'] > game_rate * 1.5:
            merchant_cancel_warning.append({
                'shopName': s['shopName'],
                'shopId': s['shopId'],
                'gameName': s['gameName'],
                'shopCancelRate': s['cancelRate'],
                'gameCancelRate': game_rate,
                'diff': s['cancelRate'] - game_rate,
                'orderCount': s['orderCount']
            })
    merchant_cancel_warning.sort(key=lambda x: x['diff'], reverse=True)

    # 渠道列表、游戏列表
    channels = sorted(set(r['channel'] for r in rows if r['channel'] and r['channel'] != '未知渠道'))
    games = sorted(set(r['gameName'] for r in rows))

    return {
        'dataType': 'sublet',
        'dateRange': date_range,
        'totals': totals,
        'channelData': channel_data,
        'xianyuDetail': xianyu_detail,
        'taobaoDetail': taobao_detail,
        'gameData': game_data,
        'top20ByAmount': top20_by_amount,
        'top20ByOrders': top20_by_orders,
        'top5AvgCancelRate': top5_avg_cancel_rate,
        'shopData': shop_data,
        'gameShopMap': game_shop_map,
        'merchantCancelWarningList': merchant_cancel_warning,
        'channels': channels,
        'games': games,
        'rawRowCount': len(rows),
        'hasChannelData': has_channel_data and len(channels) > 0
    }


# ========== 分销数据处理 ==========
def process_distribution_data(raw_rows, headers):
    f = DIST_FIELDS
    header_set = set(str(h).strip() for h in headers)
    header_idx = {h: i for i, h in enumerate(headers)}

    def get_val(row, field_name):
        idx = header_idx.get(field_name)
        if idx is not None and idx < len(row):
            return row[idx]
        return ''

    # 字段兼容
    settled_field = '正常结算订单量' if '正常结算订单量' in header_set else (
        '正常订单量' if '正常订单量' in header_set else '订单量')
    self_op_field = '自营结算订单量' if '自营结算订单量' in header_set else None
    non_self_op_field = '非自营结算订单量' if '非自营结算订单量' in header_set else None

    # 分销商ID字段兼容
    dist_id_field = f['distributorId']
    if 'userid' in header_set and f['distributorId'] in header_set:
        sample_id = ''
        sample_uid = ''
        for row in raw_rows[:100]:
            val = get_val(row, f['distributorId'])
            if val:
                sample_id = str(val)
                break
        for row in raw_rows[:100]:
            val = get_val(row, 'userid')
            if val:
                sample_uid = str(val)
                break
        import re
        if re.match(r'^\d{1,9}$', sample_id) and len(sample_uid) >= 10:
            dist_id_field = 'userid'
    elif f['distributorId'] not in header_set and 'userid' in header_set:
        dist_id_field = 'userid'

    user_pay_field = f['userPayment'] if f['userPayment'] in header_set else (
        '实付金额' if '实付金额' in header_set else f['orderAmount'])

    # 一次遍历完成所有聚合
    rows = []
    totals = {
        'orderCount': 0, 'settledOrderCount': 0, 'cancelCount': 0,
        'orderAmount': 0, 'userPayment': 0, 'platformProfit': 0,
        'distributorProfit': 0, 'distributorMarkup': 0, 'ownerProfit': 0,
        'selfOperatedOrders': 0, 'nonSelfOperatedOrders': 0
    }
    type_map = {}
    game_map = {}
    dist_map = {}  # 游戏×分销商
    dist_total_map = {}  # 分销商总计
    game_distributor_map = {}

    for row in raw_rows:
        date_val = get_val(row, f['date'])
        distributor_id = str(get_val(row, dist_id_field) or '')
        distributor_type = get_val(row, f['distributorType']) or '未知类型'
        game_name = get_val(row, f['gameName']) or '未知游戏'
        order_portal = get_val(row, f['orderPortal']) or ''
        is_self = get_val(row, f['isSelfOperated']) or ''
        order_count = safe_num(get_val(row, f['orderCount']))
        settled_order = safe_num(get_val(row, settled_field))
        cancel_count = safe_num(get_val(row, f['cancelCount']))
        order_amount = safe_num(get_val(row, f['orderAmount']))
        user_payment = safe_num(get_val(row, user_pay_field))
        platform_profit = safe_num(get_val(row, f['platformProfit']))
        platform_markup = safe_num(get_val(row, f['platformMarkup']))
        distributor_profit = safe_num(get_val(row, f['distributorProfit']))
        distributor_markup = safe_num(get_val(row, f['distributorMarkup']))
        owner_profit = safe_num(get_val(row, f['ownerProfit']))
        self_orders = safe_num(get_val(row, self_op_field)) if self_op_field else 0
        non_self_orders = safe_num(get_val(row, non_self_op_field)) if non_self_op_field else 0

        r = {
            'date': date_val,
            'distributorId': distributor_id,
            'distributorType': distributor_type,
            'gameName': game_name,
            'orderPortal': order_portal,
            'isSelfOperated': is_self,
            'orderCount': order_count,
            'settledOrderCount': settled_order,
            'cancelCount': cancel_count,
            'orderAmount': order_amount,
            'userPayment': user_payment,
            'platformProfit': platform_profit,
            'platformMarkup': platform_markup,
            'distributorProfit': distributor_profit,
            'distributorMarkup': distributor_markup,
            'ownerProfit': owner_profit,
            'selfOperatedOrders': self_orders,
            'nonSelfOperatedOrders': non_self_orders
        }
        rows.append(r)

        # totals
        totals['orderCount'] += order_count
        totals['settledOrderCount'] += settled_order
        totals['cancelCount'] += cancel_count
        totals['orderAmount'] += order_amount
        totals['userPayment'] += user_payment
        totals['platformProfit'] += platform_profit
        totals['distributorProfit'] += distributor_profit
        totals['distributorMarkup'] += distributor_markup
        totals['ownerProfit'] += owner_profit
        totals['selfOperatedOrders'] += self_orders
        totals['nonSelfOperatedOrders'] += non_self_orders

        # 类型聚合
        t = distributor_type
        if t not in type_map:
            type_map[t] = {
                'type': t, 'settledOrderCount': 0, 'platformProfit': 0,
                'settledAmount': 0, 'distributorProfit': 0, 'orderCount': 0
            }
        type_map[t]['settledOrderCount'] += settled_order
        type_map[t]['platformProfit'] += platform_profit
        type_map[t]['settledAmount'] += user_payment
        type_map[t]['distributorProfit'] += distributor_profit + distributor_markup
        type_map[t]['orderCount'] += order_count

        # 游戏聚合
        g = game_name
        if g not in game_map:
            game_map[g] = {
                'gameName': g, 'settledOrderCount': 0, 'settledAmount': 0,
                'platformProfit': 0, 'distributorProfit': 0,
                'cancelCount': 0, 'orderCount': 0
            }
        game_map[g]['settledOrderCount'] += settled_order
        game_map[g]['settledAmount'] += user_payment
        game_map[g]['platformProfit'] += platform_profit
        game_map[g]['distributorProfit'] += distributor_profit + distributor_markup
        game_map[g]['cancelCount'] += cancel_count
        game_map[g]['orderCount'] += order_count

        # 分销商×游戏聚合
        dkey = distributor_id + '_' + game_name
        if dkey not in dist_map:
            dist_map[dkey] = {
                'distributorId': distributor_id,
                'gameName': game_name,
                'distributorType': distributor_type,
                'settledOrderCount': 0, 'settledAmount': 0,
                'platformProfit': 0, 'distributorProfit': 0,
                'cancelCount': 0, 'orderCount': 0
            }
        dist_map[dkey]['settledOrderCount'] += settled_order
        dist_map[dkey]['settledAmount'] += user_payment
        dist_map[dkey]['platformProfit'] += platform_profit
        dist_map[dkey]['distributorProfit'] += distributor_profit + distributor_markup
        dist_map[dkey]['cancelCount'] += cancel_count
        dist_map[dkey]['orderCount'] += order_count

        # 分销商总计
        if distributor_id not in dist_total_map:
            dist_total_map[distributor_id] = {
                'distributorId': distributor_id,
                'distributorType': distributor_type,
                'settledOrderCount': 0, 'settledAmount': 0,
                'platformProfit': 0, 'distributorProfit': 0,
                'cancelCount': 0, 'orderCount': 0
            }
        dist_total_map[distributor_id]['settledOrderCount'] += settled_order
        dist_total_map[distributor_id]['settledAmount'] += user_payment
        dist_total_map[distributor_id]['platformProfit'] += platform_profit
        dist_total_map[distributor_id]['distributorProfit'] += distributor_profit + distributor_markup
        dist_total_map[distributor_id]['cancelCount'] += cancel_count
        dist_total_map[distributor_id]['orderCount'] += order_count

        # 游戏→分销商（用于异常预警下钻）
        if game_name not in game_distributor_map:
            game_distributor_map[game_name] = {}
        if distributor_id not in game_distributor_map[game_name]:
            game_distributor_map[game_name][distributor_id] = {
                'distributorId': distributor_id,
                'settledOrderCount': 0, 'settledAmount': 0
            }
        game_distributor_map[game_name][distributor_id]['settledOrderCount'] += settled_order
        game_distributor_map[game_name][distributor_id]['settledAmount'] += user_payment

    # 派生值
    totals['cancelRate'] = totals['cancelCount'] / totals['orderCount'] if totals['orderCount'] > 0 else 0
    totals['selfOperatedRatio'] = (
        totals['selfOperatedOrders'] / totals['settledOrderCount']
        if totals['settledOrderCount'] > 0 else 0
    )
    totals['distributorTotalProfit'] = totals['distributorProfit'] + totals['distributorMarkup']

    date_range = get_date_range(rows, 'date')

    type_data = sorted(type_map.values(), key=lambda x: x['settledOrderCount'], reverse=True)

    game_data = list(game_map.values())
    for g in game_data:
        g['cancelRate'] = g['cancelCount'] / g['orderCount'] if g['orderCount'] > 0 else 0

    top20_by_amount = sorted(game_data, key=lambda x: x['settledAmount'], reverse=True)[:20]
    top20_by_orders = sorted(game_data, key=lambda x: x['settledOrderCount'], reverse=True)[:20]

    top10_games = sorted(game_data, key=lambda x: x['settledOrderCount'], reverse=True)[:10]
    top10_avg_cancel_rate = (
        sum(g['cancelRate'] for g in top10_games) / len(top10_games)
        if top10_games else 0
    )

    distributor_game_data = list(dist_map.values())
    for d in distributor_game_data:
        d['cancelRate'] = d['cancelCount'] / d['orderCount'] if d['orderCount'] > 0 else 0

    distributor_data = list(dist_total_map.values())

    # 商户-游戏撤单率预警
    game_cancel_map = {g['gameName']: g['cancelRate'] for g in game_data}
    merchant_cancel_warning = []
    for d in distributor_game_data:
        game_rate = game_cancel_map.get(d['gameName'], 0)
        if d['orderCount'] >= 10 and d['cancelRate'] > game_rate * 1.5:
            merchant_cancel_warning.append({
                'distributorId': d['distributorId'],
                'gameName': d['gameName'],
                'merchantCancelRate': d['cancelRate'],
                'gameCancelRate': game_rate,
                'diff': d['cancelRate'] - game_rate,
                'orderCount': d['orderCount']
            })
    merchant_cancel_warning.sort(key=lambda x: x['diff'], reverse=True)

    games = sorted(set(r['gameName'] for r in rows))
    distributor_types = sorted(set(r['distributorType'] for r in rows if r['distributorType']))

    return {
        'dataType': 'distribution',
        'dateRange': date_range,
        'totals': totals,
        'typeData': type_data,
        'gameData': game_data,
        'top20ByAmount': top20_by_amount,
        'top20ByOrders': top20_by_orders,
        'top10AvgCancelRate': top10_avg_cancel_rate,
        'distributorGameData': distributor_game_data,
        'distributorData': distributor_data,
        'gameDistributorMap': game_distributor_map,
        'merchantCancelWarningList': merchant_cancel_warning,
        'games': games,
        'distributorTypes': distributor_types,
        'rawRowCount': len(rows)
    }


# ========== 环比计算 ==========
def compute_comparison(this_week, last_week, data_type):
    if not this_week or not last_week:
        return None

    result = {
        'totals': {},
        'dateRange': this_week.get('dateRange')
    }

    tw = this_week.get('totals', {})
    lw = last_week.get('totals', {})

    # 总览KPI环比
    tw_amount = tw.get('settledAmount') or tw.get('userPayment') or 0
    lw_amount = lw.get('settledAmount') or lw.get('userPayment') or 0
    result['totals']['settledAmount'] = {
        'current': tw_amount,
        'last': lw_amount,
        'change': tw_amount - lw_amount,
        'changeRate': calc_change_rate(tw_amount, lw_amount)
    }

    result['totals']['settledOrderCount'] = {
        'current': tw.get('settledOrderCount', 0),
        'last': lw.get('settledOrderCount', 0),
        'change': tw.get('settledOrderCount', 0) - lw.get('settledOrderCount', 0),
        'changeRate': calc_change_rate(tw.get('settledOrderCount', 0), lw.get('settledOrderCount', 0))
    }

    result['totals']['cancelRate'] = {
        'current': tw.get('cancelRate', 0),
        'last': lw.get('cancelRate', 0),
        'change': tw.get('cancelRate', 0) - lw.get('cancelRate', 0),
        'changePct': (tw.get('cancelRate', 0) - lw.get('cancelRate', 0)) * 100
    }

    result['totals']['platformProfit'] = {
        'current': tw.get('platformProfit', 0),
        'last': lw.get('platformProfit', 0),
        'change': tw.get('platformProfit', 0) - lw.get('platformProfit', 0),
        'changeRate': calc_change_rate(tw.get('platformProfit', 0), lw.get('platformProfit', 0))
    }

    if data_type == 'sublet':
        result['totals']['shipSuccessRate'] = {
            'current': tw.get('shipSuccessRate', 0),
            'last': lw.get('shipSuccessRate', 0),
            'change': tw.get('shipSuccessRate', 0) - lw.get('shipSuccessRate', 0)
        }
        result['totals']['monthPlatformProfit'] = tw.get('platformProfit', 0)
        result['totals']['monthOrders'] = tw.get('settledOrderCount', 0)

    if data_type == 'distribution':
        result['totals']['distributorProfit'] = {
            'current': tw.get('distributorTotalProfit', 0),
            'last': lw.get('distributorTotalProfit', 0),
            'change': tw.get('distributorTotalProfit', 0) - lw.get('distributorTotalProfit', 0),
            'changeRate': calc_change_rate(tw.get('distributorTotalProfit', 0), lw.get('distributorTotalProfit', 0))
        }
        result['totals']['selfOperatedRatio'] = tw.get('selfOperatedRatio', 0)

    # 渠道环比（转租）
    if data_type == 'sublet' and this_week.get('channelData') and last_week.get('channelData'):
        tw_has_ch = this_week.get('hasChannelData', False) and this_week.get('channels')
        lw_has_ch = last_week.get('hasChannelData', False) and last_week.get('channels')

        if tw_has_ch and lw_has_ch:
            lw_channel_map = {c['channel']: c for c in last_week['channelData']}
            result['channelComparison'] = []
            for ch in this_week['channelData']:
                last_ch = lw_channel_map.get(ch['channel'], {
                    'settledOrderCount': 0, 'platformProfit': 0, 'settledAmount': 0
                })
                result['channelComparison'].append({
                    'channel': ch['channel'],
                    'currentOrders': ch['settledOrderCount'],
                    'lastOrders': last_ch['settledOrderCount'],
                    'ordersChange': ch['settledOrderCount'] - last_ch['settledOrderCount'],
                    'ordersChangeRate': calc_change_rate(ch['settledOrderCount'], last_ch['settledOrderCount']),
                    'currentProfit': ch['platformProfit'],
                    'lastProfit': last_ch['platformProfit'],
                    'profitChange': ch['platformProfit'] - last_ch['platformProfit'],
                    'profitChangeRate': calc_change_rate(ch['platformProfit'], last_ch['platformProfit']),
                    'currentAmount': ch['settledAmount'],
                    'lastAmount': last_ch['settledAmount']
                })
        elif tw_has_ch:
            lw_total = last_week.get('totals', {})
            result['channelComparison'] = []
            for ch in this_week['channelData']:
                result['channelComparison'].append({
                    'channel': ch['channel'],
                    'currentOrders': ch['settledOrderCount'],
                    'lastOrders': 0,
                    'ordersChange': ch['settledOrderCount'],
                    'ordersChangeRate': 1,
                    'currentProfit': ch['platformProfit'],
                    'lastProfit': 0,
                    'profitChange': ch['platformProfit'],
                    'profitChangeRate': 1,
                    'currentAmount': ch['settledAmount'],
                    'lastAmount': 0
                })
            result['channelComparison'].append({
                'channel': '上周总计',
                'currentOrders': tw.get('settledOrderCount', 0),
                'lastOrders': lw.get('settledOrderCount', 0),
                'ordersChange': tw.get('settledOrderCount', 0) - lw.get('settledOrderCount', 0),
                'ordersChangeRate': calc_change_rate(tw.get('settledOrderCount', 0), lw.get('settledOrderCount', 0)),
                'currentProfit': tw.get('platformProfit', 0),
                'lastProfit': lw.get('platformProfit', 0),
                'profitChange': tw.get('platformProfit', 0) - lw.get('platformProfit', 0),
                'profitChangeRate': calc_change_rate(tw.get('platformProfit', 0), lw.get('platformProfit', 0)),
                'currentAmount': tw_amount,
                'lastAmount': lw_amount,
                'isTotalRow': True
            })
        else:
            result['channelComparison'] = [{
                'channel': '总计',
                'currentOrders': tw.get('settledOrderCount', 0),
                'lastOrders': lw.get('settledOrderCount', 0),
                'ordersChange': tw.get('settledOrderCount', 0) - lw.get('settledOrderCount', 0),
                'ordersChangeRate': calc_change_rate(tw.get('settledOrderCount', 0), lw.get('settledOrderCount', 0)),
                'currentProfit': tw.get('platformProfit', 0),
                'lastProfit': lw.get('platformProfit', 0),
                'profitChange': tw.get('platformProfit', 0) - lw.get('platformProfit', 0),
                'profitChangeRate': calc_change_rate(tw.get('platformProfit', 0), lw.get('platformProfit', 0)),
                'currentAmount': tw_amount,
                'lastAmount': lw_amount
            }]

    # 游戏环比
    if this_week.get('gameData') and last_week.get('gameData'):
        game_map = {}
        for g in this_week['gameData']:
            game_map[g['gameName']] = {'gameName': g['gameName'], 'current': g}
        for g in last_week['gameData']:
            if g['gameName'] not in game_map:
                game_map[g['gameName']] = {
                    'gameName': g['gameName'],
                    'current': {'settledOrderCount': 0, 'settledAmount': 0, 'platformProfit': 0,
                               'cancelRate': 0, 'orderCount': 0}
                }
            game_map[g['gameName']]['last'] = g

        result['gameComparison'] = []
        for item in game_map.values():
            cur = item.get('current', {})
            lst = item.get('last', {'settledOrderCount': 0, 'settledAmount': 0,
                                    'platformProfit': 0, 'cancelRate': 0, 'orderCount': 0})
            result['gameComparison'].append({
                'gameName': item['gameName'],
                'currentOrders': cur.get('settledOrderCount', 0),
                'lastOrders': lst.get('settledOrderCount', 0),
                'ordersChange': cur.get('settledOrderCount', 0) - lst.get('settledOrderCount', 0),
                'ordersChangeRate': calc_change_rate(cur.get('settledOrderCount', 0), lst.get('settledOrderCount', 0)),
                'currentAmount': cur.get('settledAmount', 0),
                'lastAmount': lst.get('settledAmount', 0),
                'amountChange': cur.get('settledAmount', 0) - lst.get('settledAmount', 0),
                'amountChangeRate': calc_change_rate(cur.get('settledAmount', 0), lst.get('settledAmount', 0)),
                'currentProfit': cur.get('platformProfit', 0),
                'lastProfit': lst.get('platformProfit', 0),
                'profitChange': cur.get('platformProfit', 0) - lst.get('platformProfit', 0),
                'profitChangeRate': calc_change_rate(cur.get('platformProfit', 0), lst.get('platformProfit', 0)),
                'currentCancelRate': cur.get('cancelRate', 0),
                'lastCancelRate': lst.get('cancelRate', 0)
            })

        # 增长TOP20
        result['growthTop20'] = sorted(
            [g for g in result['gameComparison'] if g['ordersChange'] > 0 and g['lastOrders'] > 0],
            key=lambda x: x['ordersChange'], reverse=True
        )[:20]

        # 下降TOP20
        result['declineTop20'] = sorted(
            [g for g in result['gameComparison'] if g['ordersChange'] < 0 and g['currentOrders'] > 0],
            key=lambda x: x['ordersChange']
        )[:20]

        # TOP20金额榜+环比
        result['top20AmountWithChange'] = sorted(
            result['gameComparison'], key=lambda x: x['currentAmount'], reverse=True
        )[:20]

        # TOP20订单量榜+环比
        result['top20OrdersWithChange'] = sorted(
            result['gameComparison'], key=lambda x: x['currentOrders'], reverse=True
        )[:20]

    # 店铺环比（转租）
    if data_type == 'sublet' and this_week.get('shopData') and last_week.get('shopData'):
        tw_has_ch = this_week.get('hasChannelData', False) and this_week.get('channels')
        lw_has_ch = last_week.get('hasChannelData', False) and last_week.get('channels')

        def aggregate_by_shop_game(shop_data_arr):
            agg_map = {}
            for s in shop_data_arr:
                key = shop_agg_key(s.get('sublessorId'), s.get('sublessorName'),
                                   s.get('shopId'), s.get('gameName'))
                if key not in agg_map:
                    agg_map[key] = {
                        'shopId': s.get('shopId', ''),
                        'shopName': s.get('shopName', ''),
                        'sublessorName': s.get('sublessorName', ''),
                        'sublessorId': s.get('sublessorId', ''),
                        'gameName': s.get('gameName', ''),
                        'channel': '',
                        'settledOrderCount': 0, 'settledAmount': 0,
                        'platformProfit': 0, 'sublessorProfit': 0,
                        'cancelCount': 0, 'orderCount': 0
                    }
                agg_map[key]['settledOrderCount'] += s.get('settledOrderCount', 0)
                agg_map[key]['settledAmount'] += s.get('settledAmount', 0)
                agg_map[key]['platformProfit'] += s.get('platformProfit', 0)
                agg_map[key]['sublessorProfit'] += s.get('sublessorProfit', 0)
                agg_map[key]['cancelCount'] += s.get('cancelCount', 0)
                agg_map[key]['orderCount'] += s.get('orderCount', 0)
                if s.get('shopName') and not agg_map[key]['shopName']:
                    agg_map[key]['shopName'] = s['shopName']
                if s.get('sublessorName') and not agg_map[key]['sublessorName']:
                    agg_map[key]['sublessorName'] = s['sublessorName']
            return list(agg_map.values())

        if tw_has_ch and lw_has_ch:
            tw_shop_data = this_week['shopData']
            lw_shop_data = last_week['shopData']
            use_channel_key = True
        else:
            tw_shop_data = aggregate_by_shop_game(this_week['shopData'])
            lw_shop_data = aggregate_by_shop_game(last_week['shopData'])
            use_channel_key = False

        shop_map = {}
        for s in tw_shop_data:
            key = (shop_agg_key(s.get('sublessorId'), s.get('sublessorName'),
                                s.get('shopId'), s.get('gameName'),
                                s.get('channel') if use_channel_key else None))
            shop_map[key] = {'shop': s, 'current': s}

        for s in lw_shop_data:
            key = (shop_agg_key(s.get('sublessorId'), s.get('sublessorName'),
                                s.get('shopId'), s.get('gameName'),
                                s.get('channel') if use_channel_key else None))
            if key not in shop_map:
                shop_map[key] = {
                    'shop': {**s, 'settledOrderCount': 0, 'settledAmount': 0,
                             'platformProfit': 0, 'sublessorProfit': 0,
                             'orderCount': 0, 'cancelCount': 0},
                    'current': {'settledOrderCount': 0, 'settledAmount': 0,
                                'platformProfit': 0, 'sublessorProfit': 0,
                                'orderCount': 0, 'cancelCount': 0}
                }
            shop_map[key]['last'] = s

        result['shopComparison'] = []
        for item in shop_map.values():
            cur = item.get('current', {})
            lst = item.get('last', {'settledOrderCount': 0, 'settledAmount': 0,
                                    'platformProfit': 0, 'sublessorProfit': 0,
                                    'orderCount': 0, 'cancelCount': 0})
            shop = item['shop']
            result['shopComparison'].append({
                'shopId': shop.get('shopId', ''),
                'shopName': shop.get('shopName', ''),
                'sublessorName': shop.get('sublessorName', ''),
                'gameName': shop.get('gameName', ''),
                'channel': shop.get('channel', ''),
                'currentOrders': cur.get('settledOrderCount', 0),
                'lastOrders': lst.get('settledOrderCount', 0),
                'ordersChange': cur.get('settledOrderCount', 0) - lst.get('settledOrderCount', 0),
                'ordersChangeRate': calc_change_rate(cur.get('settledOrderCount', 0), lst.get('settledOrderCount', 0)),
                'currentAmount': cur.get('settledAmount', 0),
                'lastAmount': lst.get('settledAmount', 0),
                'amountChange': cur.get('settledAmount', 0) - lst.get('settledAmount', 0),
                'amountChangeRate': calc_change_rate(cur.get('settledAmount', 0), lst.get('settledAmount', 0)),
                'currentProfit': cur.get('platformProfit', 0),
                'lastProfit': lst.get('platformProfit', 0),
                'profitChange': cur.get('platformProfit', 0) - lst.get('platformProfit', 0),
                'profitChangeRate': calc_change_rate(cur.get('platformProfit', 0), lst.get('platformProfit', 0))
            })

    # 分销商环比（分销）
    if data_type == 'distribution' and this_week.get('distributorData') and last_week.get('distributorData'):
        dist_map = {}
        for d in this_week['distributorData']:
            dist_map[d['distributorId']] = {'dist': d, 'current': d}
        for d in last_week['distributorData']:
            if d['distributorId'] not in dist_map:
                dist_map[d['distributorId']] = {
                    'dist': {**d, 'settledOrderCount': 0, 'settledAmount': 0,
                             'platformProfit': 0, 'distributorProfit': 0,
                             'orderCount': 0, 'cancelCount': 0},
                    'current': {'settledOrderCount': 0, 'settledAmount': 0,
                                'platformProfit': 0, 'distributorProfit': 0,
                                'orderCount': 0, 'cancelCount': 0}
                }
            dist_map[d['distributorId']]['last'] = d

        result['distributorComparison'] = []
        for item in dist_map.values():
            cur = item.get('current', {})
            lst = item.get('last', {'settledOrderCount': 0, 'settledAmount': 0,
                                    'platformProfit': 0, 'distributorProfit': 0,
                                    'orderCount': 0, 'cancelCount': 0})
            dist = item['dist']
            result['distributorComparison'].append({
                'distributorId': dist.get('distributorId', ''),
                'distributorType': dist.get('distributorType', ''),
                'currentOrders': cur.get('settledOrderCount', 0),
                'lastOrders': lst.get('settledOrderCount', 0),
                'ordersChange': cur.get('settledOrderCount', 0) - lst.get('settledOrderCount', 0),
                'ordersChangeRate': calc_change_rate(cur.get('settledOrderCount', 0), lst.get('settledOrderCount', 0)),
                'currentAmount': cur.get('settledAmount', 0),
                'lastAmount': lst.get('settledAmount', 0),
                'amountChange': cur.get('settledAmount', 0) - lst.get('settledAmount', 0),
                'amountChangeRate': calc_change_rate(cur.get('settledAmount', 0), lst.get('settledAmount', 0)),
                'currentProfit': cur.get('platformProfit', 0),
                'lastProfit': lst.get('platformProfit', 0),
                'profitChange': cur.get('platformProfit', 0) - lst.get('platformProfit', 0),
                'profitChangeRate': calc_change_rate(cur.get('platformProfit', 0), lst.get('platformProfit', 0))
            })

    # 游戏×转租商交叉分析（转租）
    if data_type == 'sublet' and this_week.get('shopData'):
        result['subletCrossAnalysis'] = {}

        # 上周数据聚合
        last_week_shop_map = {}
        last_week_sublessor_map = {}
        if last_week and last_week.get('shopData'):
            for s in last_week['shopData']:
                shop_key = shop_agg_key(s.get('sublessorId'), s.get('sublessorName'),
                                        s.get('shopId'), s.get('gameName'))
                sub_key = shop_agg_key(s.get('sublessorId'), s.get('sublessorName'),
                                       '', s.get('gameName'))

                if shop_key not in last_week_shop_map:
                    last_week_shop_map[shop_key] = {
                        'settledOrderCount': 0, 'settledAmount': 0,
                        'platformProfit': 0, 'sublessorProfit': 0
                    }
                last_week_shop_map[shop_key]['settledOrderCount'] += s.get('settledOrderCount', 0)
                last_week_shop_map[shop_key]['settledAmount'] += s.get('settledAmount', 0)
                last_week_shop_map[shop_key]['platformProfit'] += s.get('platformProfit', 0)
                last_week_shop_map[shop_key]['sublessorProfit'] += s.get('sublessorProfit', 0)

                if sub_key not in last_week_sublessor_map:
                    last_week_sublessor_map[sub_key] = {
                        'settledOrderCount': 0, 'settledAmount': 0,
                        'platformProfit': 0, 'sublessorProfit': 0
                    }
                last_week_sublessor_map[sub_key]['settledOrderCount'] += s.get('settledOrderCount', 0)
                last_week_sublessor_map[sub_key]['settledAmount'] += s.get('settledAmount', 0)
                last_week_sublessor_map[sub_key]['platformProfit'] += s.get('platformProfit', 0)
                last_week_sublessor_map[sub_key]['sublessorProfit'] += s.get('sublessorProfit', 0)

        # 本周数据聚合
        this_week_shop_map = {}
        this_week_sublessor_map = {}

        for s in this_week['shopData']:
            shop_key = shop_agg_key(s.get('sublessorId'), s.get('sublessorName'),
                                    s.get('shopId'), s.get('gameName'))
            sub_key = shop_agg_key(s.get('sublessorId'), s.get('sublessorName'),
                                   '', s.get('gameName'))

            if shop_key not in this_week_shop_map:
                this_week_shop_map[shop_key] = {
                    'shopId': s.get('shopId', ''), 'shopName': s.get('shopName', ''),
                    'sublessorName': s.get('sublessorName', ''),
                    'sublessorId': s.get('sublessorId', ''),
                    'gameName': s.get('gameName', ''),
                    'channel': s.get('channel', ''),
                    'settledAmount': 0, 'settledOrderCount': 0,
                    'sublessorProfit': 0, 'platformProfit': 0,
                    'orderCount': 0, 'cancelCount': 0
                }
            this_week_shop_map[shop_key]['settledAmount'] += s.get('settledAmount', 0)
            this_week_shop_map[shop_key]['settledOrderCount'] += s.get('settledOrderCount', 0)
            this_week_shop_map[shop_key]['sublessorProfit'] += s.get('sublessorProfit', 0)
            this_week_shop_map[shop_key]['platformProfit'] += s.get('platformProfit', 0)
            this_week_shop_map[shop_key]['orderCount'] += s.get('orderCount', 0)
            this_week_shop_map[shop_key]['cancelCount'] += s.get('cancelCount', 0)
            if s.get('shopName') and not this_week_shop_map[shop_key]['shopName']:
                this_week_shop_map[shop_key]['shopName'] = s['shopName']
            if s.get('sublessorName') and not this_week_shop_map[shop_key]['sublessorName']:
                this_week_shop_map[shop_key]['sublessorName'] = s['sublessorName']
            if s.get('channel') and not this_week_shop_map[shop_key]['channel']:
                this_week_shop_map[shop_key]['channel'] = s['channel']

            if sub_key not in this_week_sublessor_map:
                this_week_sublessor_map[sub_key] = {
                    'sublessorId': s.get('sublessorId', ''),
                    'sublessorName': s.get('sublessorName', ''),
                    'gameName': s.get('gameName', ''),
                    'settledAmount': 0, 'settledOrderCount': 0,
                    'sublessorProfit': 0, 'platformProfit': 0,
                    'orderCount': 0, 'cancelCount': 0,
                    'shopIds': set()
                }
            this_week_sublessor_map[sub_key]['settledAmount'] += s.get('settledAmount', 0)
            this_week_sublessor_map[sub_key]['settledOrderCount'] += s.get('settledOrderCount', 0)
            this_week_sublessor_map[sub_key]['sublessorProfit'] += s.get('sublessorProfit', 0)
            this_week_sublessor_map[sub_key]['platformProfit'] += s.get('platformProfit', 0)
            this_week_sublessor_map[sub_key]['orderCount'] += s.get('orderCount', 0)
            this_week_sublessor_map[sub_key]['cancelCount'] += s.get('cancelCount', 0)
            this_week_sublessor_map[sub_key]['shopIds'].add(s.get('shopId') or '')
            if s.get('sublessorName') and not this_week_sublessor_map[sub_key]['sublessorName']:
                this_week_sublessor_map[sub_key]['sublessorName'] = s['sublessorName']

        # 对每个游戏构建转租商列表
        for game in this_week.get('games', []):
            sublessors = []
            for s in this_week_sublessor_map.values():
                if s['gameName'] != game:
                    continue
                last_key = shop_agg_key(s['sublessorId'], s['sublessorName'], '', game)
                last = last_week_sublessor_map.get(last_key, {
                    'settledOrderCount': 0, 'settledAmount': 0,
                    'platformProfit': 0, 'sublessorProfit': 0
                })
                sublessors.append({
                    'sublessorId': s['sublessorId'],
                    'sublessorName': s['sublessorName'] or '未知转租商',
                    'settledAmount': s['settledAmount'],
                    'settledOrderCount': s['settledOrderCount'],
                    'sublessorProfit': s['sublessorProfit'],
                    'platformProfit': s['platformProfit'],
                    'lastSettledAmount': last.get('settledAmount', 0),
                    'lastSettledOrderCount': last.get('settledOrderCount', 0),
                    'lastSublessorProfit': last.get('sublessorProfit', 0),
                    'lastPlatformProfit': last.get('platformProfit', 0),
                    'amountChange': s['settledAmount'] - last.get('settledAmount', 0),
                    'amountChangeRate': calc_change_rate(s['settledAmount'], last.get('settledAmount', 0)),
                    'ordersChange': s['settledOrderCount'] - last.get('settledOrderCount', 0),
                    'ordersChangeRate': calc_change_rate(s['settledOrderCount'], last.get('settledOrderCount', 0)),
                    'shopCount': len(s['shopIds']),
                    'shops': []
                })

            sublessors.sort(key=lambda x: x['settledAmount'], reverse=True)

            # 分级
            total = len(sublessors)
            for idx, sub in enumerate(sublessors):
                pct = (idx + 1) / total if total > 0 else 0
                if pct <= 0.2:
                    sub['grade'] = 'S'
                elif pct <= 0.4:
                    sub['grade'] = 'A'
                elif pct <= 0.6:
                    sub['grade'] = 'B'
                elif pct <= 0.8:
                    sub['grade'] = 'C'
                else:
                    sub['grade'] = 'D'

            # 为每个转租商填充店铺列表
            for sub in sublessors:
                sub_shops = []
                for s in this_week_shop_map.values():
                    match = False
                    if sub['sublessorId'] and s.get('sublessorId'):
                        match = s['gameName'] == game and s['sublessorId'] == sub['sublessorId']
                    else:
                        match = s['gameName'] == game and s.get('sublessorName') == sub['sublessorName']
                    if not match:
                        continue

                    last_shop_key = shop_agg_key(s.get('sublessorId'), s.get('sublessorName'),
                                                 s.get('shopId'), game)
                    last = last_week_shop_map.get(last_shop_key, {
                        'settledOrderCount': 0, 'settledAmount': 0,
                        'platformProfit': 0, 'sublessorProfit': 0
                    })
                    sub_shops.append({
                        'shopId': s.get('shopId', ''),
                        'shopName': s.get('shopName', ''),
                        'channel': s.get('channel', ''),
                        'settledAmount': s.get('settledAmount', 0),
                        'settledOrderCount': s.get('settledOrderCount', 0),
                        'platformProfit': s.get('platformProfit', 0),
                        'lastSettledAmount': last.get('settledAmount', 0),
                        'lastSettledOrderCount': last.get('settledOrderCount', 0),
                        'lastPlatformProfit': last.get('platformProfit', 0),
                        'amountChange': s.get('settledAmount', 0) - last.get('settledAmount', 0),
                        'amountChangeRate': calc_change_rate(s.get('settledAmount', 0), last.get('settledAmount', 0)),
                        'ordersChange': s.get('settledOrderCount', 0) - last.get('settledOrderCount', 0),
                        'ordersChangeRate': calc_change_rate(s.get('settledOrderCount', 0), last.get('settledOrderCount', 0))
                    })
                sub_shops.sort(key=lambda x: x['settledAmount'], reverse=True)
                sub['shops'] = sub_shops

            result['subletCrossAnalysis'][game] = sublessors

        # TOP10游戏
        result['subletTop10Games'] = [
            g['gameName'] for g in sorted(
                this_week.get('gameData', []),
                key=lambda x: x['settledAmount'], reverse=True
            )[:10]
        ]

    # 游戏×分销商交叉分析（分销）
    if data_type == 'distribution' and this_week.get('distributorGameData'):
        result['crossAnalysis'] = {}
        for game in this_week.get('games', []):
            result['crossAnalysis'][game] = [
                {
                    'distributorId': d['distributorId'],
                    'amount': d['settledAmount'],
                    'orders': d['settledOrderCount'],
                    'distributorProfit': d['distributorProfit'],
                    'platformProfit': d['platformProfit'],
                    'cancelRate': d['cancelRate']
                }
                for d in this_week['distributorGameData']
                if d['gameName'] == game
            ]

    return result


# ========== Excel 读取 ==========
def read_sheet(xlsx_path, use_readonly=True):
    wb = load_workbook(xlsx_path, read_only=use_readonly, data_only=True)
    ws = wb[wb.sheetnames[0]]
    sheet_name = wb.sheetnames[0]

    headers = []
    rows = []

    for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
        if row_idx == 1:
            headers = [str(cell) if cell is not None else '' for cell in row]
        else:
            row_data = []
            for cell in row:
                if cell is None:
                    row_data.append('')
                elif isinstance(cell, (int, float)):
                    row_data.append(cell)
                else:
                    row_data.append(str(cell))
            rows.append(row_data)

    wb.close()
    return headers, rows, sheet_name


# ========== 单文件转换 ==========
def convert_xlsx_to_json(xlsx_path, generate_summary=True):
    """将 xlsx 文件转换为 JSON 格式（原始明细 + 预计算摘要）"""
    if not os.path.exists(xlsx_path):
        print(f"错误: 文件不存在 - {xlsx_path}", file=sys.stderr)
        return False

    file_size = os.path.getsize(xlsx_path)
    print(f"正在处理: {os.path.basename(xlsx_path)} ({file_size/1024/1024:.2f}MB)")

    try:
        # 读取 Excel
        print("  读取 Excel...")
        try:
            headers, rows, sheet_name = read_sheet(xlsx_path, True)
            if file_size > 50 * 1024 and (len(rows) < 10 or len(headers) <= 1):
                print(f"  read_only模式数据不完整，回退到普通模式...")
                headers, rows, sheet_name = read_sheet(xlsx_path, False)
        except Exception:
            headers, rows, sheet_name = read_sheet(xlsx_path, False)

        print(f"  数据量: {len(rows)}行 × {len(headers)}列")

        # 构建原始明细 JSON（紧凑格式）
        result = {
            "sheetName": sheet_name,
            "headers": headers,
            "rows": rows
        }

        json_path = os.path.splitext(xlsx_path)[0] + '.json'
        with open(json_path, 'w', encoding='utf-8') as f:
            json.dump(result, f, ensure_ascii=False, separators=(',', ':'))
        json_size = os.path.getsize(json_path)
        print(f"  原始明细JSON: {os.path.basename(json_path)} ({json_size/1024/1024:.2f}MB)")

        # 生成预计算摘要
        if generate_summary:
            data_type = detect_data_type(headers)
            print(f"  数据类型: {data_type}")

            if data_type == 'sublet':
                summary = process_sublet_data(rows, headers)
            elif data_type == 'distribution':
                summary = process_distribution_data(rows, headers)
            else:
                print("  警告: 无法识别数据类型，跳过摘要生成")
                return True

            summary_path = os.path.splitext(xlsx_path)[0] + '_summary.json'
            with open(summary_path, 'w', encoding='utf-8') as f:
                json.dump(summary, f, ensure_ascii=False, separators=(',', ':'))
            summary_size = os.path.getsize(summary_path)
            print(f"  预计算摘要: {os.path.basename(summary_path)} ({summary_size/1024:.2f}KB)")
            print(f"  压缩比: {json_size/summary_size:.1f}x  (原始:{json_size/1024/1024:.2f}MB → 摘要:{summary_size/1024:.2f}KB)")

        return True

    except Exception as e:
        print(f"转换失败: {e}", file=sys.stderr)
        import traceback
        traceback.print_exc()
        return False


# ========== 批处理模式 ==========
def batch_process(data_dir):
    """批量处理目录下所有 xlsx 文件，并生成对比摘要"""
    if not os.path.isdir(data_dir):
        print(f"错误: 目录不存在 - {data_dir}", file=sys.stderr)
        return False

    print(f"批量处理目录: {data_dir}")
    print("=" * 60)

    # 查找所有 xlsx 文件
    xlsx_files = {}
    for fname in os.listdir(data_dir):
        if fname.lower().endswith('.xlsx') and not fname.startswith('~$'):
            fpath = os.path.join(data_dir, fname)
            name_no_ext = os.path.splitext(fname)[0]
            xlsx_files[name_no_ext] = fpath

    if not xlsx_files:
        print("未找到 xlsx 文件")
        return False

    print(f"找到 {len(xlsx_files)} 个 xlsx 文件")

    # 逐个转换
    results = {}
    for name, fpath in sorted(xlsx_files.items()):
        success = convert_xlsx_to_json(fpath, generate_summary=True)
        results[name] = success
        print()

    # 生成对比摘要
    print("=" * 60)
    print("生成对比摘要...")

    # 转租对比
    sublet_current = None
    sublet_previous = None
    for prefix in ['sublet_current', 'sublet_prev', 'sublet_last']:
        for name in xlsx_files:
            if name == prefix:
                summary_path = os.path.join(data_dir, name + '_summary.json')
                if os.path.exists(summary_path):
                    with open(summary_path, 'r', encoding='utf-8') as f:
                        data = json.load(f)
                    if prefix == 'sublet_current':
                        sublet_current = data
                    else:
                        sublet_previous = data
                    break

    # 尝试 alternative names
    if sublet_previous is None:
        for name in xlsx_files:
            if 'sublet' in name and ('previous' in name or 'prev' in name or 'last' in name):
                summary_path = os.path.join(data_dir, name + '_summary.json')
                if os.path.exists(summary_path):
                    with open(summary_path, 'r', encoding='utf-8') as f:
                        sublet_previous = json.load(f)
                    break

    if sublet_current and sublet_previous:
        comparison = compute_comparison(sublet_current, sublet_previous, 'sublet')
        comp_path = os.path.join(data_dir, 'sublet_comparison.json')
        with open(comp_path, 'w', encoding='utf-8') as f:
            json.dump(comparison, f, ensure_ascii=False, separators=(',', ':'))
        comp_size = os.path.getsize(comp_path)
        print(f"  转租对比摘要: {os.path.basename(comp_path)} ({comp_size/1024:.2f}KB)")
    else:
        print(f"  转租对比: 缺少数据 (current={sublet_current is not None}, previous={sublet_previous is not None})")

    # 分销对比
    dist_current = None
    dist_previous = None
    for name in xlsx_files:
        if name == 'dist_current':
            summary_path = os.path.join(data_dir, name + '_summary.json')
            if os.path.exists(summary_path):
                with open(summary_path, 'r', encoding='utf-8') as f:
                    dist_current = json.load(f)
        elif 'dist' in name and ('previous' in name or 'prev' in name or 'last' in name):
            summary_path = os.path.join(data_dir, name + '_summary.json')
            if os.path.exists(summary_path):
                with open(summary_path, 'r', encoding='utf-8') as f:
                    dist_previous = json.load(f)

    if dist_current and dist_previous:
        comparison = compute_comparison(dist_current, dist_previous, 'distribution')
        comp_path = os.path.join(data_dir, 'dist_comparison.json')
        with open(comp_path, 'w', encoding='utf-8') as f:
            json.dump(comparison, f, ensure_ascii=False, separators=(',', ':'))
        comp_size = os.path.getsize(comp_path)
        print(f"  分销对比摘要: {os.path.basename(comp_path)} ({comp_size/1024:.2f}KB)")
    else:
        print(f"  分销对比: 缺少数据 (current={dist_current is not None}, previous={dist_previous is not None})")

    print()
    print("批量处理完成!")
    return True


# ========== 主函数 ==========
def main():
    if len(sys.argv) < 2:
        print("用法:")
        print("  python convert_xlsx_to_json.py <xlsx文件路径>")
        print("  python convert_xlsx_to_json.py --batch <data目录路径>")
        print()
        print("说明:")
        print("  单文件模式: 转换单个xlsx，生成原始JSON + 预计算摘要")
        print("  批处理模式: 转换目录下所有xlsx，并生成对比摘要")
        sys.exit(1)

    if sys.argv[1] == '--batch':
        if len(sys.argv) < 3:
            print("错误: 请指定目录路径", file=sys.stderr)
            sys.exit(1)
        data_dir = sys.argv[2]
        if not os.path.isabs(data_dir):
            data_dir = os.path.abspath(data_dir)
        success = batch_process(data_dir)
        sys.exit(0 if success else 1)
    else:
        xlsx_path = sys.argv[1]
        if not os.path.isabs(xlsx_path):
            xlsx_path = os.path.abspath(xlsx_path)
        success = convert_xlsx_to_json(xlsx_path)
        sys.exit(0 if success else 1)


if __name__ == '__main__':
    main()
